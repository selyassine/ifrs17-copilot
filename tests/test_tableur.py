"""Classeurs Excel et CSV (D-079) : détection, conversion en tableau, pièges, chaîne complète."""

import io
import json

import pytest

pytest.importorskip("openpyxl")
from openpyxl import Workbook  # noqa: E402

from multimodal.spreadsheet import CSV_MIME, XLSX_MIME, csv_text, is_xls, xlsx_text  # noqa: E402
from multimodal.vision import detect_mime  # noqa: E402

TRUTH = {"group_id": "AUTO-X", "currency": "EUR", "unit_multiplier": 1000, "premiums": [100.0, 0.0],
         "outflows": [45.0, 50.0], "coverage_units": [1000, 900], "discount_rates": 0.03, "risk_adjustment_pct": 0.05}


def _xlsx(**kw) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Flux"
    rows = [["Groupe", "AUTO-X"], ["Devise", "EUR"], ["Montants en", "milliers"], ["Taux", 0.03], ["Ajustement", 0.05],
            [], ["Période", "Primes", "Sinistres", "Frais", "Contrats"],
            ["Année 1", 100, 40, 5, 1000], ["Année 2", 0, 45.5, 4.5, 900]]
    for r in rows:
        ws.append(r)
    ws["B4"].number_format = ws["B5"].number_format = "0.00%"
    if kw.get("hidden"):
        h = wb.create_sheet("Brouillon")
        h.append(["Primes", 999])
        h.sheet_state = "hidden"
    if kw.get("formula"):
        ws["B8"] = "=60+40"
    if kw.get("text_number"):
        ws["C9"] = "45,50"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_detection_par_le_contenu():
    assert detect_mime(_xlsx()) == XLSX_MIME
    assert detect_mime(b"Periode;Primes;Sinistres\nAnnee 1;100,00;40,00\nAnnee 2;0,00;45,50\n") == CSV_MIME
    assert detect_mime(b"a,b\n1,2\n", "flux.csv") == CSV_MIME
    assert detect_mime(b"Bonjour, ceci est un texte.") is None
    assert detect_mime(b"%PDF-1.7 ...") == "application/pdf"
    assert is_xls(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 10)


def test_classeur_converti_comme_un_tableau():
    d = xlsx_text(_xlsx())
    assert "Taux | 3 %" in d.text and "Ajustement | 5 %" in d.text
    assert "Année 2 | 0 | 45,5 | 4,5 | 900" in d.text
    assert 0.03 in d.table_numbers and 45.5 in d.table_numbers
    assert "feuille Flux, cellule C9" in d.where(45.5)
    assert not d.problems and not d.notes


def test_pieges_du_tableur():
    d = xlsx_text(_xlsx(hidden=True, formula=True, text_number=True))
    assert any("feuilles cachées ignorées : Brouillon" in n for n in d.notes) and "999" not in d.text
    assert any("formule sans valeur" in p and "B8" in p for p in d.problems)
    assert 45.5 in d.table_numbers                      # « 45,50 » saisi comme du texte


def test_csv_a_la_francaise_et_a_l_anglaise():
    d = csv_text("Période;Primes;Sinistres\nAnnée 1;1 234,50;40,00\nAnnée 2;0,00;45,50\n".encode("cp1252"))
    assert "Année 1 | 1 234,50 | 40,00" in d.text and 1234.5 in d.table_numbers
    d = csv_text(b"Period,Premiums,Claims\nYear 1,100.5,40\n")
    assert 100.5 in d.table_numbers


def _fake(answer):
    from core.config import Settings
    from core.llm import LLMClient
    return LLMClient(Settings(llm_provider="fake", llm_max_retries=0), fake_responder=lambda m, mo: json.dumps(answer))


def test_extraction_d_un_classeur():
    from multimodal.extract import extract_from_text
    ans = TRUTH | {"manquants": [], "groupes_detectes": ["AUTO-X"]}
    r = extract_from_text(xlsx_text(_xlsx()), _fake(ans), mode="tableur")
    assert r.status == "ok" and r.mode == "tableur" and "cellule" in r.provenance["Primes"]
    r = extract_from_text(xlsx_text(_xlsx(formula=True)), _fake(ans), mode="tableur")
    assert r.status == "a_verifier" and any("formule sans valeur" in i for i in r.issues)
    r = extract_from_text(xlsx_text(_xlsx(hidden=True)), _fake(ans), mode="tableur")
    assert r.status == "ok" and any("cachées" in w for w in r.warnings)


def test_api_accepte_excel_et_refuse_xls(tmp_path, monkeypatch):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient

    import api.main as main
    from api.main import app
    from core.config import Settings, get_settings
    s = Settings(llm_provider="fake", api_token="secret", trace_path=tmp_path / "t.jsonl", registry_path=tmp_path / "r.db")
    app.dependency_overrides[get_settings] = lambda: s
    monkeypatch.setattr(main, "make_client", lambda st: _fake(TRUTH | {"manquants": [], "groupes_detectes": ["AUTO-X"]}))
    monkeypatch.setattr(main, "_retriever_loader", lambda st: None)
    try:
        c = TestClient(app)
        r = c.post("/v1/case", headers={"X-API-Token": "secret"},
                   files={"file": ("flux.xlsx", _xlsx(), XLSX_MIME)})
        assert r.status_code == 200, r.text
        assert r.json()["status"] == "note_prete" and "lecture des cellules du classeur" in r.json()["note_html"]
        r = c.post("/v1/case", headers={"X-API-Token": "secret"},
                   files={"file": ("vieux.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, "application/vnd.ms-excel")})
        assert r.status_code == 415 and ".xlsx" in r.json()["detail"]
    finally:
        app.dependency_overrides.clear()


def test_jeu_excel(tmp_path):
    from multimodal.excel_cases import generate
    path = generate(tmp_path / "excel")
    cases = [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]
    assert len(cases) == 22 and sum(c["attendu"] == "alerte" for c in cases) == 3
    assert {c["layout"] for c in cases} >= {"lignes", "colonnes", "deux_feuilles", "csv", "feuille_cachee"}


def test_n8n_et_site_acceptent_les_tableurs():
    from core.config import ROOT
    wf = (ROOT / "workflows" / "ifrs17_copilot.json").read_text(encoding="utf-8")
    assert "spreadsheetml.sheet" in wf and ".xlsx, .csv" in wf
    js = (ROOT / "api" / "web" / "assets" / "js" / "views" / "intake.js").read_text(encoding="utf-8")
    assert ".xlsx,.csv" in js


def test_vraisemblance_et_pourcentage_aberrant(tmp_path):
    """D-080 : un chiffre présent dans le document mais aberrant (montant au format %, période × 80) est signalé."""
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    from agents.schemas import PortfolioData
    from multimodal.extract import plausibility_issues
    from multimodal.spreadsheet import XLSX_MIME, spreadsheet_text

    ok = PortfolioData(group_id="G", premiums=[1000, 0, 0, 0], outflows=[300, 250, 200, 150],
                       coverage_units=[10, 9, 8, 7], discount_rates=[0.02, 0.025, 0.03, 0.03], risk_adjustment_pct=0.05)
    assert plausibility_issues(ok) == []
    assert plausibility_issues(ok.model_copy(update={"outflows": [300, 250, 200, 15000]}))
    assert plausibility_issues(ok.model_copy(update={"discount_rates": [2.0, 2.5, 3.0, 3.0]}))
    assert plausibility_issues(ok.model_copy(update={"risk_adjustment_pct": 0.8}))
    wb = Workbook()
    ws = wb.active
    ws.append(["Sinistres", 300.0, 250.0, 64241.07])
    ws["D1"].number_format = "0.00%"
    path = tmp_path / "x.xlsx"
    wb.save(path)
    doc = spreadsheet_text(path.read_bytes(), XLSX_MIME)
    assert any("pourcentage invraisemblable" in p and "D1" in p for p in doc.problems)


def test_generateur_excel_formate_la_ligne_des_taux_en_colonnes(tmp_path):
    """Le format % suit la colonne d'origine des taux, même transposée (bogue du premier jeu : année 6 en %)."""
    pytest.importorskip("openpyxl")
    from multimodal.excel_cases import generate
    from multimodal.spreadsheet import XLSX_MIME, spreadsheet_text
    generate(tmp_path)
    doc = spreadsheet_text((tmp_path / "xl10.xlsx").read_bytes(), XLSX_MIME)
    assert not doc.problems and "Taux spot | 2,25 %" in doc.text and "64241,07" in doc.text
