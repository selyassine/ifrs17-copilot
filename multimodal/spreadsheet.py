"""Lecture des classeurs Excel (.xlsx) et des fichiers CSV (D-079).

Chez un assureur, les flux d'un portefeuille arrivent souvent en tableur. Un classeur est le format le plus sûr :
les nombres sont exacts (ni OCR ni colonnes mal alignées). Il est converti dans le même format que la couche
texte d'un PDF (``DocumentText`` : lignes de tableau avec des cellules séparées par « | »), si bien que toute
la chaîne reste identique : LLM, contrôle d'ancrage, calcul, note. Chaque nombre est localisé à sa cellule
(« feuille Flux, cellule C4 ») pour la provenance.

Pièges propres aux tableurs, traités par le code :
- pourcentage stocké 0,03 et affiché 3 % : la cellule est écrite « 3 % » (ce que voit l'utilisateur) ;
- formule sans valeur enregistrée (fichier produit par un programme, jamais ouvert dans Excel) : la valeur est
  inconnue → problème bloquant, validation humaine ;
- feuilles cachées : ignorées et signalées (une donnée cachée ne doit pas entrer dans un calcul sans être vue) ;
- nombres stockés comme du texte (« 1 234,50 ») : lus comme tout texte ;
- classeur Excel 97-2003 (.xls) : refusé avec un message clair (enregistrer en .xlsx).
"""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime

from multimodal.documents import MAX_CHARS, DocumentText, _cell_number, numbers_in_text

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv"
MAX_SHEETS = 5
MAX_ROWS = 200
MAX_COLS = 30


class SpreadsheetError(ValueError):
    """Classeur illisible ou format non pris en charge."""


def is_xlsx(data: bytes) -> bool:
    if not data.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return "xl/workbook.xml" in z.namelist()
    except zipfile.BadZipFile:
        return False


def is_xls(data: bytes) -> bool:
    return data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"      # conteneur OLE (Excel 97-2003)


def _decode(data: bytes) -> str | None:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return None


def looks_like_csv(data: bytes, filename: str = "") -> bool:
    """Texte avec un séparateur régulier (« ; », « , » ou tabulation) et des chiffres."""
    if filename.lower().endswith((".csv", ".tsv")):
        return _decode(data[:20000]) is not None
    text = _decode(data[:20000])
    if text is None or "\x00" in text or not any(ch.isdigit() for ch in text):
        return False
    lines = [ln for ln in text.splitlines() if ln.strip()][:20]
    if len(lines) < 2:
        return False
    for sep in (";", "\t", ","):
        counts = [ln.count(sep) for ln in lines]
        if min(counts) >= 1 and max(counts) - min(counts) <= 1:
            return True
    return False


def _fmt_number(v: float) -> str:
    if float(v).is_integer():
        return str(int(v))
    return f"{v:.10g}".replace(".", ",")


def _cell_text(value, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "VRAI" if value else "FAUX"
    if isinstance(value, datetime | date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, int | float):
        if "%" in (number_format or ""):
            return f"{_fmt_number(round(value * 100, 10))} %"
        return _fmt_number(value)
    return str(value).replace("\n", " ").replace("|", "/").strip()


def _column(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _build(sheets: list[tuple[str, list[list[tuple[str, str]]]]], notes: list[str], problems: list[str],
           truncated: bool) -> DocumentText:
    """``sheets`` : (nom, lignes) ; une ligne = cellules (texte affiché, coordonnée)."""
    parts, table_nums, text_nums, locations = [], [], [], {}
    for name, rows in sheets:
        lines = []
        for row in rows:
            cells = [c for c, _ in row]
            while cells and not cells[-1]:
                cells.pop()
            if not any(cells):
                continue
            lines.append(" | ".join(cells))
            for text, coord in row:
                nums = _cell_number(text)
                if nums:
                    table_nums += nums
                else:
                    # une phrase dans une cellule (« Taux spot : 5,59 % ; 5,74 % … ») : ses nombres comptent comme
                    # du texte courant (trouvé par la mesure LIC du 24/10, classeur lc06 : fausse alerte)
                    nums = numbers_in_text(text)
                    text_nums += nums
                for v in nums:
                    for key in {round(v, 4), round(v, 2)}:
                        locations.setdefault(key, set()).add(f"feuille {name}, cellule {coord}")
        if lines:
            parts.append(f"[Feuille « {name} »]\n[Tableau 1, feuille {name}]\n" + "\n".join(lines))
    if notes:
        parts.append("[Remarques du lecteur de classeur]\n" + "\n".join(notes))
    text = "\n\n".join(parts)
    doc = DocumentText(text=text[:MAX_CHARS], table_numbers=table_nums, text_numbers=text_nums,
                       pages=len(sheets), truncated=truncated or len(text) > MAX_CHARS, locations=locations)
    doc.notes = notes
    doc.problems = problems
    return doc


def xlsx_text(data: bytes) -> DocumentText:
    from openpyxl import load_workbook

    try:
        values = load_workbook(io.BytesIO(data), data_only=True)
        formulas = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as exc:            # classeur corrompu ou chiffré
        raise SpreadsheetError(f"classeur illisible : {exc}"[:300]) from exc
    sheets, notes, problems = [], [], []
    truncated = False
    visible = [ws for ws in values.worksheets if ws.sheet_state == "visible"]
    hidden = [ws.title for ws in values.worksheets if ws.sheet_state != "visible"]
    if hidden:
        notes.append(f"feuilles cachées ignorées : {', '.join(hidden)}")
    if len(visible) > MAX_SHEETS:
        truncated = True
        notes.append(f"seules les {MAX_SHEETS} premières feuilles visibles ont été lues ({len(visible)} au total)")
    for ws in visible[:MAX_SHEETS]:
        wf = formulas[ws.title]
        rows = []
        if ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS:
            truncated = True
        for r in range(1, min(ws.max_row, MAX_ROWS) + 1):
            row = []
            for c in range(1, min(ws.max_column, MAX_COLS) + 1):
                cell = ws.cell(row=r, column=c)
                coord = f"{_column(c)}{r}"
                raw = wf.cell(row=r, column=c).value
                if cell.value is None and isinstance(raw, str) and raw.startswith("="):
                    problems.append(f"formule sans valeur enregistrée : feuille {ws.title}, cellule {coord} "
                                    "(ouvrir et enregistrer le classeur dans Excel)")
                    row.append(("?", coord))
                    continue
                if (isinstance(cell.value, int | float) and not isinstance(cell.value, bool)
                        and "%" in (cell.number_format or "") and abs(cell.value) > 10):
                    # « 6 424 107 % » : un montant au format pourcentage (D-080) ; Excel l'affiche ainsi, on le signale
                    shown = f"{cell.value * 100:,.0f}".replace(",", " ")
                    problems.append(f"pourcentage invraisemblable ({shown} %) : feuille {ws.title}, cellule {coord} "
                                    "(montant au format pourcentage ?)")
                row.append((_cell_text(cell.value, cell.number_format), coord))
            rows.append(row)
        sheets.append((ws.title, rows))
    if not sheets:
        raise SpreadsheetError("classeur sans feuille visible")
    return _build(sheets, notes, problems, truncated)


def csv_text(data: bytes) -> DocumentText:
    text = _decode(data)
    if text is None:
        raise SpreadsheetError("fichier CSV illisible (encodage inconnu)")
    sample = "\n".join(text.splitlines()[:20])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        sep = dialect.delimiter
    except csv.Error:
        sep = ";" if sample.count(";") >= sample.count(",") else ","
    rows = []
    truncated = False
    for r, rec in enumerate(csv.reader(io.StringIO(text), delimiter=sep), start=1):
        if r > MAX_ROWS:
            truncated = True
            break
        rows.append([(cell.replace("|", "/").strip(), f"{_column(c)}{r}") for c, cell in enumerate(rec[:MAX_COLS], start=1)])
    return _build([("CSV", rows)], [], [], truncated)


def spreadsheet_text(data: bytes, mime: str) -> DocumentText:
    return xlsx_text(data) if mime == XLSX_MIME else csv_text(data)
